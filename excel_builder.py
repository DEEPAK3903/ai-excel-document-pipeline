import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import logging

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")

def build_excel_report(structured_data, output_filename="processed_document_report.xlsx"):
    """
    Converts structured JSON data into a professional multi-sheet Excel report.
    Sheet 1: Invoice Summary
    Sheet 2: Line Items Breakdown
    """
    try:
        # Sheet 1 Data
        summary_data = {
            "Field": ["Invoice Number", "Date", "Customer Name", "Subtotal ($)", "Tax ($)", "Total Amount ($)"],
            "Value": [
                structured_data.get("invoice_number"),
                structured_data.get("date"),
                structured_data.get("customer_name"),
                structured_data.get("subtotal"),
                structured_data.get("tax"),
                structured_data.get("total_amount")
            ]
        }
        df_summary = pd.DataFrame(summary_data)

        # Sheet 2 Data
        line_items = structured_data.get("line_items", [])
        df_items = pd.DataFrame(line_items)

        # Create Excel Writer
        with pd.ExcelWriter(output_filename, engine='openpyxl') as writer:
            df_summary.to_excel(writer, sheet_name="Invoice Summary", index=False)
            if not df_items.empty:
                df_items.to_excel(writer, sheet_name="Line Items", index=False)

        # Apply OpenPyXL Styles
        wb = openpyxl.load_workbook(output_filename)
        header_fill = PatternFill(start_color="142D73", end_color="142D73", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
                ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 28

        wb.save(output_filename)
        logging.info(f"Excel report generated successfully: {output_filename}")
        return output_filename
    except Exception as e:
        logging.error(f"Failed to generate Excel report: {e}")
        return None
